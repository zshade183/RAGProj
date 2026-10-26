"""Extract searchable table, chart and native connector relationships from PPTX."""

from __future__ import annotations

from html import escape
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE


def clean_text(value: Any) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\v", "\n").strip()


def escape_markdown_cell(value: str) -> str:
    return clean_text(value).replace("|", "\\|").replace("\n", "<br>")


def rows_to_markdown(rows: list[list[str]]) -> str:
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = "| " + " | ".join(escape_markdown_cell(cell) for cell in normalized[0]) + " |"
    divider = "| " + " | ".join("---" for _ in range(width)) + " |"
    body = [
        "| " + " | ".join(escape_markdown_cell(cell) for cell in row) + " |"
        for row in normalized[1:]
    ]
    return "\n".join([header, divider, *body])


def rows_to_html(rows: list[list[str]]) -> str:
    cells = []
    for index, row in enumerate(rows):
        tag = "th" if index == 0 else "td"
        cells.append(
            "<tr>" + "".join(f"<{tag}>{escape(cell)}</{tag}>" for cell in row) + "</tr>"
        )
    return "<table>" + "".join(cells) + "</table>"


def extract_table(shape, base: dict) -> dict:
    rows, cells = [], []
    for row_index, row in enumerate(shape.table.rows):
        values = []
        for column_index, cell in enumerate(row.cells):
            text = "" if cell.is_spanned else clean_text(cell.text)
            values.append(text)
            cells.append(
                {
                    "row": row_index,
                    "column": column_index,
                    "text": clean_text(cell.text),
                    "is_merge_origin": bool(cell.is_merge_origin),
                    "is_spanned": bool(cell.is_spanned),
                    "row_span": int(cell.span_height),
                    "column_span": int(cell.span_width),
                }
            )
        rows.append(values)
    markdown = rows_to_markdown(rows)
    table = {
        "rows": rows,
        "cells": cells,
        "row_count": len(rows),
        "column_count": len(shape.table.columns),
        "markdown": markdown,
    }
    return {
        **base,
        "kind": "table",
        "text": markdown,
        "table": table,
        "rows": rows,
        "cells": cells,
        "row_count": table["row_count"],
        "column_count": table["column_count"],
        "text_as_html": rows_to_html(rows),
    }


def extract_chart(shape, base: dict) -> dict:
    chart = shape.chart
    title = ""
    try:
        if chart.has_title:
            title = clean_text(chart.chart_title.text_frame.text)
    except (AttributeError, ValueError):
        title = ""
    categories = []
    try:
        categories = [clean_text(value) for value in chart.plots[0].categories]
    except (AttributeError, IndexError, KeyError, TypeError, ValueError):
        pass
    series = []
    for item in chart.series:
        try:
            values = [value for value in item.values]
        except (AttributeError, TypeError):
            values = []
        series.append({"name": clean_text(item.name), "values": values})
    lines = ["[图表]"]
    if title:
        lines.append(f"标题：{title}")
    lines.append(f"类型：{chart.chart_type}")
    if categories:
        lines.append("类别：" + "、".join(categories))
    for item in series:
        values = "、".join("" if value is None else str(value) for value in item["values"])
        lines.append(f"系列 {item['name'] or '未命名'}：{values}")
    return {
        **base,
        "kind": "chart",
        "text": "\n".join(lines),
        "chart": {
            "chart_type": str(chart.chart_type),
            "title": title,
            "categories": categories,
            "series": series,
        },
    }


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)


def _shape_label(shape) -> str:
    if getattr(shape, "has_text_frame", False):
        value = clean_text(shape.text_frame.text)
        if value:
            return value
    return clean_text(getattr(shape, "name", ""))


def extract_connector_relations(shapes) -> list[dict]:
    all_shapes = list(walk_shapes(shapes))
    by_id = {str(shape.shape_id): shape for shape in all_shapes}
    relations = []
    for connector in all_shapes:
        begin = connector.element.xpath(".//*[local-name()='stCxn']")
        end = connector.element.xpath(".//*[local-name()='endCxn']")
        if not begin and not end:
            continue
        begin_id = begin[0].get("id") if begin else ""
        end_id = end[0].get("id") if end else ""
        source = by_id.get(str(begin_id))
        target = by_id.get(str(end_id))
        source_label = _shape_label(source) if source is not None else f"shape-{begin_id or '?'}"
        target_label = _shape_label(target) if target is not None else f"shape-{end_id or '?'}"
        label = _shape_label(connector)
        if not source_label and not target_label:
            continue
        text = f"[关系] {source_label or '起点'} -[{label or '连接'}]-> {target_label or '终点'}"
        relations.append(
            {
                "source": "native",
                "kind": "relation",
                "text": text,
                "include_in_retrieval": True,
                "relation": {
                    "source_shape_id": str(begin_id) if begin_id else None,
                    "target_shape_id": str(end_id) if end_id else None,
                    "source_label": source_label,
                    "target_label": target_label,
                    "label": label,
                    "connector_name": clean_text(getattr(connector, "name", "")),
                },
            }
        )
    return relations
