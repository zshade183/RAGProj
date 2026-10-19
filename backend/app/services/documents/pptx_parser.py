from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def normalize_text(text: str) -> str:
    """保留换行，将 PowerPoint 软换行转换为普通换行。"""

    return text.replace("\r\n", "\n").replace("\v", "\n").strip()


def parse_shape(shape, shape_path: str) -> dict[str, Any]:
    """读取一个形状；组合形状递归处理。"""

    block: dict[str, Any] = {
        "shape_path": shape_path,
        "shape_id": shape.shape_id,
        "name": shape.name,
        "shape_type": str(shape.shape_type),
        "bbox_emu": [
            int(shape.left),
            int(shape.top),
            int(shape.width),
            int(shape.height),
        ],
    }

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        block.update(
            {
                "type": "group",
                "children": [
                    parse_shape(child, f"{shape_path}/{child.shape_id}") for child in shape.shapes
                ],
            }
        )
        return block

    if shape.has_table:
        cells = []

        for row_index, row in enumerate(shape.table.rows):
            for column_index, cell in enumerate(row.cells):
                cells.append(
                    {
                        "row": row_index,
                        "column": column_index,
                        "text": normalize_text(cell.text),
                        "is_merge_origin": cell.is_merge_origin,
                        "is_spanned": cell.is_spanned,
                        "row_span": cell.span_height,
                        "column_span": cell.span_width,
                    }
                )

        block.update(
            {
                "type": "table",
                "row_count": len(shape.table.rows),
                "column_count": len(shape.table.columns),
                "cells": cells,
            }
        )
        return block

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        block.update(
            {
                "type": "image",
                "requires_followup": True,
                "reason": "图片文字及图示尚未识别",
            }
        )
        return block

    if shape.has_chart:
        block.update(
            {
                "type": "chart",
                "requires_followup": True,
                "reason": "图表数据及视觉关系尚未提取",
            }
        )
        return block

    if shape.has_text_frame:
        paragraphs = [
            {
                "text": normalize_text(paragraph.text),
                "level": paragraph.level,
            }
            for paragraph in shape.text_frame.paragraphs
        ]

        block.update(
            {
                "type": "text",
                "text": "\n".join(
                    paragraph["text"] for paragraph in paragraphs if paragraph["text"]
                ),
                "paragraphs": paragraphs,
            }
        )

        # 不把公式误认为普通文本已完整提取。
        has_math = bool(
            shape.element.xpath(".//*[local-name()='oMath' or local-name()='oMathPara']")
        )

        if has_math:
            block["requires_followup"] = True
            block["reason"] = "包含公式，需要补充公式解析或页面视觉识别"

        return block

    block.update(
        {
            "type": "other",
            "requires_followup": True,
            "reason": "未提取的图形或嵌入内容，需要页面检查",
        }
    )
    return block


def walk_blocks(blocks):
    """遍历普通块以及组合形状中的所有子块。"""

    for block in blocks:
        yield block
        if block["type"] == "group":
            yield from walk_blocks(block["children"])


def parse_pptx(file_path: Path) -> dict[str, Any]:
    """提取 PPTX 原生元素，不执行 OCR。"""

    presentation = Presentation(str(file_path))
    pages = []

    for page_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = normalize_text(title_shape.text) if title_shape else ""

        blocks = [parse_shape(shape, str(shape.shape_id)) for shape in slide.shapes]

        notes = ""
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes = normalize_text(notes_frame.text)

        pages.append(
            {
                "page_number": page_number,
                "title": title,
                "notes": notes,
                "blocks": blocks,
            }
        )

    all_blocks = [block for page in pages for block in walk_blocks(page["blocks"])]

    text_characters = sum(
        len(block.get("text", "")) for block in all_blocks if block["type"] == "text"
    )

    return {
        "schema_version": "1.0",
        "parser": "python-pptx-native",
        "extraction_stage": "native_only",
        "slide_width_emu": int(presentation.slide_width),
        "slide_height_emu": int(presentation.slide_height),
        "page_count": len(pages),
        "summary": {
            "text_characters": text_characters,
            "table_count": sum(block["type"] == "table" for block in all_blocks),
            "image_count": sum(block["type"] == "image" for block in all_blocks),
            "followup_block_count": sum(
                block.get("requires_followup", False) for block in all_blocks
            ),
        },
        "limitations": [
            "元素顺序为文件中的形状顺序，不保证符合视觉阅读顺序",
            "组合形状内部坐标未转换成幻灯片绝对坐标",
            "未进行图片 OCR、图表解读及完整公式解析",
            "未展开母版、布局和背景中的全部可见内容",
        ],
        "pages": pages,
    }
