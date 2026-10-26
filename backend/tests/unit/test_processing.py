import hashlib
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from app.tasks import queue
from app.tasks.pipeline import make_page, parse, text_document
from fastapi import HTTPException


@pytest.mark.asyncio
async def test_duplicate_enqueue_reuses_running_job():
    doc = SimpleNamespace(id=uuid4(), status="parsing")
    job = SimpleNamespace(status="running")
    session = AsyncMock()
    session.scalar.return_value = job
    assert await queue.enqueue(session, doc) is job
    session.flush.assert_not_awaited()


@pytest.mark.asyncio
@pytest.mark.parametrize("status", ["uploading", "deleting"])
async def test_enqueue_rejects_invalid_state(status):
    with pytest.raises(HTTPException) as exc:
        await queue.enqueue(AsyncMock(), SimpleNamespace(status=status))
    assert exc.value.status_code == 409


@pytest.mark.asyncio
async def test_retry_preserves_active_snapshot():
    doc = SimpleNamespace(
        id=uuid4(), status="failed", active_snapshot_id="old", error_message="failure"
    )
    job = SimpleNamespace(status="failed", attempts=3)
    session = AsyncMock()
    session.scalar.return_value = job
    await queue.enqueue(session, doc)
    assert job.status == "queued" and job.attempts == 0
    assert doc.active_snapshot_id == "old"
    assert doc.error_message is None


@pytest.mark.asyncio
@pytest.mark.parametrize("token_matches,expired", [(False, False), (True, True), (True, False)])
async def test_lease_fencing(token_matches, expired):
    now = datetime.now(UTC)
    token = uuid4()
    task = {"document_id": str(uuid4()), "job_id": str(uuid4()), "token": str(token)}
    doc = SimpleNamespace(status="indexing")
    job = SimpleNamespace(
        status="running",
        token=token if token_matches else uuid4(),
        lease_until=now + timedelta(seconds=-1 if expired else 60),
    )
    session = AsyncMock()
    session.scalar.side_effect = [doc, job, now]
    locked_doc, locked_job = await queue.lock_owned(session, task)
    assert (locked_job is job) == (token_matches and not expired)
    assert (locked_doc is doc) == (token_matches and not expired)


@pytest.mark.parametrize(
    "suffix,text,expected",
    [
        (".txt", "密码学基础\nRSA", "RSA"),
        (".md", "# 标题\n正文", "标题"),
        (".html", "<script>BAD</script><p>正文</p>", "正文"),
    ],
)
def test_text_formats(tmp_path, suffix, text, expected):
    if suffix == ".md":
        import nltk
        try:
            nltk.data.find("tokenizers/punkt_tab")
            nltk.data.find("taggers/averaged_perceptron_tagger_eng")
        except LookupError:
            pytest.skip("Unstructured Markdown integration requires downloaded NLTK data")
    source = tmp_path / ("input" + suffix)
    source.write_text(text, encoding="utf-8")
    task = {"filename": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest()}
    doc = parse(task, source, tmp_path, None, lambda *a, **k: None)
    assert doc["stage"] == "merged"
    assert expected in doc["pages"][0]["blocks"][0]["text"]
    assert "BAD" not in doc["pages"][0]["blocks"][0]["text"]


def test_unsupported_format_fails(tmp_path):
    source = tmp_path / "input.exe"
    source.write_bytes(b"bad")
    with pytest.raises(ValueError, match="不支持"):
        parse({"filename": source.name}, source, tmp_path, None, lambda *a, **k: None)


def test_chunk_quality_exclusion():
    from app.services.documents.chunking import build_chunks

    from scripts.filter_ocr_quality import apply_policy

    doc = text_document(
        {"filename": "x.txt", "sha256": "a" * 64}, [make_page(1, "这是完整原生文字且需要保留")]
    )
    doc["pages"][0]["blocks"].append(
        {
            "block_id": "ocr",
            "source": "ocr",
            "kind": "text",
            "text": "这是完整原生文字且需要保留",
            "include_in_retrieval": True,
            "words": [],
        }
    )
    filtered = apply_policy(doc)
    chunks = build_chunks(filtered, uuid4(), uuid4())
    assert len(chunks) == 1
    assert len(chunks[0]["source_blocks"]) == 1


def test_table_structure_survives_chunking():
    from app.services.documents.chunking import build_chunks
    from app.services.documents.visual_structure import rows_to_markdown

    rows = [["步骤", "操作"], ["1", "字节代替"], ["2", "行移位"]]
    markdown = rows_to_markdown(rows)
    document = text_document({"filename": "table.md", "sha256": "a" * 64}, [make_page(1, markdown)])
    document["pages"][0]["blocks"][0].update(
        {
            "kind": "table",
            "table": {
                "rows": rows,
                "row_count": 3,
                "column_count": 2,
                "markdown": markdown,
            },
        }
    )
    chunks = build_chunks(document, uuid4(), uuid4())
    assert "| 步骤 | 操作 |" in chunks[0]["text"]
    assert chunks[0]["source_blocks"][0]["structure"]["table"]["rows"] == rows


def test_pptx_table_chart_and_native_connector_extraction(tmp_path):
    from app.services.documents.visual_structure import (
        extract_chart,
        extract_connector_relations,
        extract_table,
    )
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    table_shape.table.cell(0, 0).text = "步骤"
    table_shape.table.cell(0, 1).text = "操作"
    table_shape.table.cell(1, 0).text = "1"
    table_shape.table.cell(1, 1).text = "字节代替"
    table = extract_table(table_shape, {"shape_id": table_shape.shape_id})
    assert table["table"]["rows"] == [["步骤", "操作"], ["1", "字节代替"]]
    assert "| 步骤 | 操作 |" in table["text"]

    data = ChartData()
    data.categories = ["A", "B"]
    data.add_series("数值", (1, 2))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(3),
        Inches(4),
        Inches(3),
        data,
    )
    chart = extract_chart(chart_shape, {"shape_id": chart_shape.shape_id})
    assert chart["chart"]["categories"] == ["A", "B"]
    assert chart["chart"]["series"][0]["values"] == [1, 2]

    source = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(1), Inches(0.5))
    source.text = "字节代替"
    target = slide.shapes.add_textbox(Inches(3), Inches(6), Inches(1), Inches(0.5))
    target.text = "行移位"
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(2),
        Inches(6.25),
        Inches(3),
        Inches(6.25),
    )
    connector.begin_connect(source, 1)
    connector.end_connect(target, 3)
    relations = extract_connector_relations(slide.shapes)
    assert any(
        relation["relation"]["source_label"] == "字节代替"
        and relation["relation"]["target_label"] == "行移位"
        for relation in relations
    )


def test_actual_course_pptx_exposes_structured_tables_and_relations():
    from pathlib import Path

    from app.services.documents.visual_structure import extract_connector_relations
    from pptx import Presentation

    from scripts.merge_pptx_ocr import extract_shapes

    path = (
        Path(__file__).resolve().parents[3]
        / "docs"
        / "网络认证技术-02+密码学基础.pptx"
    )
    if not path.is_file():
        pytest.skip("Actual PPTX fixture not present")
    blocks = []
    for slide in Presentation(str(path)).slides:
        blocks.extend(extract_shapes(slide.shapes))
        blocks.extend(extract_connector_relations(slide.shapes))
    assert sum(block.get("kind") == "table" for block in blocks) >= 1
    assert any(block.get("kind") == "relation" for block in blocks)
    assert any(block.get("table", {}).get("rows") for block in blocks)


@pytest.mark.parametrize("failure", [None, "count", "artifact"])
def test_pipeline_result_requires_index_and_artifacts(tmp_path, monkeypatch, failure):
    import json
    import sys
    from types import ModuleType

    import numpy as np
    from app.core import config
    from app.services.retrieval import hybrid_store
    from app.services.storage import minio_storage
    from app.tasks import pipeline

    raw = "对称加密使用共享密钥。".encode()
    task = {
        "filename": "test.txt",
        "sha256": hashlib.sha256(raw).hexdigest(),
        "document_id": str(uuid4()),
        "knowledge_base_id": str(uuid4()),
        "token": str(uuid4()),
        "object_name": "test.txt",
        "bucket_name": "documents",
    }
    folder = tmp_path / task["token"]
    folder.mkdir()
    uploaded = []

    class Storage:
        def download_file(self, **kwargs):
            kwargs["file_path"].write_bytes(raw)

        def upload_file(self, **kwargs):
            if failure == "artifact":
                raise OSError("injected storage outage")
            uploaded.append(kwargs["object_name"])

    class Embedder:
        max_tokens = 512

        def __init__(self, *args, **kwargs):
            pass

        def token_count(self, text):
            return len(text) + 2

        def encode_documents(self, texts, **kwargs):
            vectors = np.zeros((len(texts), 1024), dtype=np.float32)
            vectors[:, 0] = 1
            return vectors

    class Store:
        def __init__(self, **kwargs):
            self.rows = []

        def ensure_collection(self):
            pass

        def upsert(self, rows):
            self.rows = rows

        def count_snapshot(self, *args):
            return 0 if failure == "count" else len(self.rows)

        def close(self):
            pass

    fake_module = ModuleType("app.services.embeddings.bge")
    fake_module.MODEL_NAME = "BAAI/bge-large-zh-v1.5"
    fake_module.QUERY_INSTRUCTION = "为这个句子生成表示以用于检索相关文章："
    fake_module.BgeEmbedding = Embedder
    monkeypatch.setitem(sys.modules, "app.services.embeddings.bge", fake_module)
    monkeypatch.setattr(pipeline, "ROOT", tmp_path)
    monkeypatch.setattr(
        config,
        "get_settings",
        lambda: SimpleNamespace(worker_embedding_device="cpu", milvus_uri="unused", rag_semantic_chunking=False),
    )
    monkeypatch.setattr(minio_storage, "get_storage", Storage)
    monkeypatch.setattr(hybrid_store, "HybridStore", Store)
    if failure:
        with pytest.raises((RuntimeError, OSError)):
            pipeline.run(task, folder)
        assert not (folder / "result.json").exists()
    else:
        pipeline.run(task, folder)
        result = json.loads((folder / "result.json").read_text())
        assert pipeline.digest(tmp_path / result["manifest_path"]) == result["snapshot_id"]
        assert any(path.endswith("embedding-manifest.json") for path in uploaded)
