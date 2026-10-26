import json
from copy import deepcopy
from uuid import uuid4

import httpx
import numpy as np
import pytest
from app.services.rag.answering import AnsweringService
from app.services.rag.context import expand_context
from app.services.rag.decisions import Grade, Judge, WebSearch
from app.services.rag.query_planning import looks_compound, merge_query_hits
from app.services.rag.reference_graphs import build_reference_graph1
from app.services.rag.workflow import build_workflow, initial_state
from app.services.retrieval.hybrid_store import HybridStore
from langchain_core.embeddings import Embeddings
from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
from langchain_core.messages import AIMessage
from langgraph.checkpoint.memory import InMemorySaver


def fixture_snapshot():
    kb, doc = str(uuid4()), str(uuid4())
    chunks = [
        {
            "chunk_id": str(uuid4()),
            "document_id": doc,
            "knowledge_base_id": kb,
            "source_filename": "course.pptx",
            "page_number": p,
            "text": t,
            "source_blocks": [],
        }
        for p, t in [(43, "字节代替"), (44, "行移位"), (45, "列混淆")]
    ]
    return {"knowledge_base_id": kb, "document_id": doc, "chunks": chunks}, "a" * 64


def hit(c, **extra):
    return {"distance": 0.8, "entity": deepcopy(c), **extra}


def test_neighbors_boundaries_and_budget():
    m, sid = fixture_snapshot()
    evidence = expand_context([hit(m["chunks"][0])], [(m, sid)])
    assert [h["entity"]["page_number"] for h in evidence] == [43, 44]
    assert evidence[1]["distance"] is None
    assert evidence[1]["retrieval_kind"] == "adjacent"
    assert len(expand_context([hit(m["chunks"][0])], [(m, sid)], max_sources=1)) == 1
    assert not expand_context([], [(m, sid)])
    assert len(expand_context([hit(m["chunks"][0])], [(m, sid)], max_chars=4)) == 1


def test_context_rejects_old_versions_and_foreign_docs():
    m, sid = fixture_snapshot()
    old = hit({**m["chunks"][0], "snapshot_id": "b" * 64})
    with pytest.raises(ValueError):
        expand_context([old], [(m, sid)])
    foreign = hit({**m["chunks"][0], "document_id": str(uuid4())})
    with pytest.raises(ValueError):
        expand_context([foreign], [(m, sid)])


def test_context_does_not_cross_documents_and_omits_bad_ocr():
    m, sid = fixture_snapshot()
    other, _ = fixture_snapshot()
    m["chunks"][1]["text"] = "行移位GARBAGE"
    m["chunks"][1]["source_blocks"] = [
        {"source": "ocr", "needs_review": True, "chunk_char_start": 3, "chunk_char_end": 10}
    ]
    evidence = expand_context([hit(m["chunks"][0])], [(m, sid), (other, sid)])
    assert len(evidence) == 2
    assert evidence[1]["entity"]["text"] == "行移位"


class Answerer:
    def answer(self, question, hits, subquestions=None):
        return {
            "answer": "行移位 [S1]",
            "citation_ids": ["S1"],
            "sources": [{"text": "行移位"}],
            "citation_check": "valid_ids",
        }

    def complete(self, *a, **kw):
        return "你好"


class FakeJudge:
    def __init__(self, route="retrieve", relevant=True, grounded=True, useful=True):
        self.target, self.rel, self.ground, self.use = route, relevant, grounded, useful
        self.histories = []

    def route(self, *a):
        return self.target

    def rewrite(self, question, query, history=None):
        self.histories.append(history)
        return question + " 改写"

    def relevant(self, *a):
        return self.rel

    def grounded(self, *a):
        return self.ground

    def useful(self, *a):
        return self.use


class FakeWeb:
    def __init__(self):
        self.calls = 0

    def search(self, *a):
        self.calls += 1
        return [{"entity": {"text": "web"}}]


def run_graph(judge=None, web=None, checkpointer=None, empty=False):
    return build_workflow(
        lambda q: ([] if empty else [{"entity": {"text": "x"}}], 20),
        lambda h: h,
        Answerer(),
        judge or FakeJudge(),
        web or FakeWeb(),
        checkpointer,
    )


def test_adaptive_success():
    s = run_graph().invoke(initial_state("问题", "adaptive", False))
    assert s["result"]["status"] == "answered"
    assert [t["node"] for t in s["trace"]] == [
        "route",
        "retrieve",
        "grade_documents",
        "generate",
        "verify",
        "finish",
    ]


def test_rewrites_bounded_no_web_permission():
    web = FakeWeb()
    s = run_graph(web=web, empty=True).invoke(initial_state("问题", "adaptive", False))
    assert s["rewrites"] == 2 and web.calls == 0
    assert s["result"]["status"] == "insufficient_evidence"


def test_web_after_rewrites_and_direct_web_denied():
    web = FakeWeb()
    s = run_graph(web=web, empty=True).invoke(initial_state("问题", "adaptive", True))
    assert s["rewrites"] == 2 and web.calls == 1
    web = FakeWeb()
    s = run_graph(judge=FakeJudge(route="web"), web=web).invoke(
        initial_state("问题", "adaptive", False)
    )
    assert web.calls == 0 and s["result"]["status"] == "answered"


def test_hallucination_never_published_after_limit():
    s = run_graph(FakeJudge(grounded=False)).invoke(initial_state("问题", "adaptive", False))
    assert s["generations"] == 3
    assert s["result"]["status"] == "insufficient_evidence"
    assert s["result"]["sources"] == []


def test_unhelpful_answer_rewrites_and_stops():
    s = run_graph(FakeJudge(useful=False)).invoke(initial_state("问题", "adaptive", False))
    assert s["rewrites"] == 2 and s["generations"] == 3
    assert s["result"]["status"] == "insufficient_evidence"


def test_agent_mode_memory_is_thread_scoped():
    judge = FakeJudge()
    graph = run_graph(judge, checkpointer=InMemorySaver())
    cfg = {"configurable": {"thread_id": "one"}}
    graph.invoke(initial_state("第一问", "agent", False), cfg)
    s = graph.invoke(initial_state("第二问", "agent", False), cfg)
    assert len(s["history"]) == 2 and judge.histories[-1][0]["question"] == "第一问"
    assert "verify" not in [t["node"] for t in s["trace"]]
    other = graph.invoke(
        initial_state("新会话", "agent", False), {"configurable": {"thread_id": "two"}}
    )
    assert len(other["history"]) == 1


def test_structured_judge_fails_closed():
    class Invalid:
        def complete(self, *a, **kw):
            return '{"binary_score":"maybe"}'

    with pytest.raises(RuntimeError):
        Judge(Invalid()).structured("test", {}, Grade)


def test_answer_prompt_preserves_page_url_and_rejects_fake_citation():
    service = AnsweringService("https://example.invalid", "test", "model")
    requests = []

    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(
            200, json={"choices": [{"message": {"content": "答案 [S9]"}, "finish_reason": "stop"}]}
        )

    service.client.close()
    service.client = httpx.Client(transport=httpx.MockTransport(respond))
    try:
        m, _ = fixture_snapshot()
        with pytest.raises(RuntimeError, match="不存在"):
            service.answer("问题", [hit(m["chunks"][0])])
        data = json.loads(requests[0]["messages"][1]["content"])
        assert data["evidence"][0]["page_number"] == 43
    finally:
        service.close()


def test_hybrid_both_paths_have_same_scope():
    class Client:
        def hybrid_search(self, **kw):
            self.kw = kw
            return [[]]

    store = HybridStore.__new__(HybridStore)
    store.client = Client()
    kb, doc = str(uuid4()), str(uuid4())
    store.hybrid_search("字节代替", [0.0] * 1024, [(kb, doc, "a" * 64)])
    requests = store.client.kw["reqs"]
    assert requests[0].expr == requests[1].expr
    assert doc in requests[0].expr and "a" * 64 in requests[0].expr
    assert {r.anns_field for r in requests} == {"vector", "sparse"}


def test_web_requires_key():
    with pytest.raises(RuntimeError):
        WebSearch("").search("test")


class TinyEmbedder:
    max_tokens = 80

    def token_count(self, text):
        return len(text)

    def encode_documents(self, texts, **kwargs):
        return np.array([[1.0, (sum(map(ord, t)) % 19) / 19] for t in texts])


def test_semantic_token_budget_and_provenance():
    from app.services.documents.semantic import semantic_chunks
    from app.tasks.pipeline import make_page, text_document

    doc = text_document(
        {"filename": "x.txt", "sha256": "a" * 64},
        [make_page(1, "第一句介绍加密。第二句说明密钥。\n" * 30)],
    )
    chunks = semantic_chunks(doc, uuid4(), uuid4(), TinyEmbedder())
    assert all(len(c["text"]) <= 80 for c in chunks)
    assert all(c["source_blocks"] for c in chunks)
    assert len({c["chunk_id"] for c in chunks}) == len(chunks)
    assert all(c["page_number"] == 1 for c in chunks)


def test_markdown_hierarchy_from_unstructured_elements(monkeypatch, tmp_path):
    from app.services.documents.semantic import markdown_page
    from langchain_community.document_loaders import UnstructuredMarkdownLoader
    from langchain_core.documents import Document

    docs = [
        Document(page_content="章", metadata={"category": "Title", "element_id": "a"}),
        Document(
            page_content="节", metadata={"category": "Title", "element_id": "b", "parent_id": "a"}
        ),
        Document(page_content="内容", metadata={"category": "NarrativeText", "parent_id": "b"}),
    ]
    monkeypatch.setattr(UnstructuredMarkdownLoader, "load", lambda self: docs)
    result = markdown_page(tmp_path / "x.md")
    assert result["blocks"][-1]["heading_path"] == ["章", "节"]
    assert result["blocks"][-1]["text"] == "内容"


def test_actual_course_native_neighbor():
    from pathlib import Path

    from app.services.documents.chunking import build_chunks
    from app.tasks.pipeline import make_page, text_document
    from pptx import Presentation

    courses = list((Path(__file__).resolve().parents[4] / "upload").glob("*.pptx"))
    # External fixture isn't shipped with a code-only update.
    if not courses:
        pytest.skip("Original PPTX fixture not present")
    presentation = Presentation(courses[0])
    pages = [
        make_page(i + 1, "\n".join(s.text for s in slide.shapes if s.has_text_frame))
        for i, slide in enumerate(presentation.slides)
    ]
    kb, doc = uuid4(), uuid4()
    document = text_document({"filename": courses[0].name, "sha256": "a" * 64}, pages)
    chunks = build_chunks(document, kb, doc)
    manifest = {"knowledge_base_id": str(kb), "document_id": str(doc), "chunks": chunks}
    seed = next(c for c in chunks if c["page_number"] == 43 and "字节代替" in c["text"])
    evidence = expand_context([hit(seed)], [(manifest, "a" * 64)])
    assert any(
        e["entity"]["page_number"] == 44 and "循环左移" in e["entity"]["text"] for e in evidence
    )


def test_runtime_hybrid_conversation_cleanup(monkeypatch):
    from app.services.rag import api_service, decisions
    from app.services.retrieval import hybrid_store
    m, sid = fixture_snapshot()
    m["collection_name"] = hybrid_store.HYBRID_COLLECTION
    m["chunk_count"] = 3
    class Store:
        def __init__(self, **kw):
            pass
        def ensure_collection(self):
            pass
        def count_snapshot(self, *a):
            return 3
        def hybrid_search(self, *a):
            return [hit(m["chunks"][0])]
        def close(self):
            pass
    class Embed:
        def encode_query(self, q):
            return np.zeros(1024)
    class Rank:
        def rank(self, question, hits, **kw):
            return [{**h, "rerank_score": 1} for h in hits]
    runtime = api_service.RagRuntime.__new__(api_service.RagRuntime)
    runtime.answering, runtime.embedder, runtime.reranker = Answerer(), Embed(), Rank()
    monkeypatch.setattr(hybrid_store, "HybridStore", Store)
    judge = FakeJudge()
    monkeypatch.setattr(decisions, "Judge", lambda a: judge)
    session = uuid4()
    a = runtime.ask(m, sid, "字节代替", 20, 5, .3, session_id=session)
    b = runtime.ask(m, sid, "它呢", 20, 5, .3, session_id=session)
    assert a["status"] == b["status"] == "answered"
    assert b["retrieval_mode"] == "hybrid_rrf"
    assert judge.histories[-1][0]["question"] == "字节代替"
    assert len(runtime.memories) == 1


def test_reference_graph1_bind_tools_and_toolnode():
    m, _ = fixture_snapshot()
    class ToolCallingFake(FakeMessagesListChatModel):
        def bind_tools(self, tools, **kwargs):
            return self

    model = ToolCallingFake(
        responses=[
            AIMessage(
                content="",
                tool_calls=[
                    {
                        "name": "rag_retriever",
                        "args": {"query": "行移位"},
                        "id": "call-1",
                        "type": "tool_call",
                    }
                ],
            )
        ]
    )
    graph = build_reference_graph1(
        llm=model,
        retrieve=lambda query: ([hit(m["chunks"][1])], 1, [query]),
        expand=lambda hits: hits,
        answering=Answerer(),
        judge=FakeJudge(),
        max_rewrites=2,
    )
    state = graph.invoke(
        {
            "messages": [("user", "行移位如何处理？")],
            "question": "行移位如何处理？",
            "evidence": [],
            "accepted_count": 0,
            "retrieved_count": 0,
            "rewrites": 0,
            "trace": [],
        }
    )
    assert state["result"]["answer"] == "行移位 [S1]"
    assert [item["node"] for item in state["trace"]] == [
        "agent",
        "retrieve",
        "grade_documents",
        "generate",
    ]


def test_compound_query_detection_and_retrieval_fusion():
    assert looks_compound("请比较 EUV 与 TSV，并分别说明")
    assert not looks_compound("行移位是什么")
    m, _ = fixture_snapshot()
    first = {
        **hit(m["chunks"][0]),
        "rerank_score": 0.4,
    }
    stronger = {
        **hit(m["chunks"][0]),
        "rerank_score": 0.9,
    }
    second = {
        **hit(m["chunks"][1]),
        "rerank_score": 0.7,
    }
    selected, total = merge_query_hits([[first, second], [stronger]], top_k=2, min_score=0.3)
    assert total == 3
    assert [item["rerank_score"] for item in selected] == [0.9, 0.7]


def test_judge_decomposes_compound_question():
    class Planner:
        def complete(self, *args, **kwargs):
            return '{"queries":["EUV 光刻的作用","TSV 的作用"]}'

    assert Judge(Planner()).decompose("比较 EUV 和 TSV") == [
        "EUV 光刻的作用",
        "TSV 的作用",
    ]


def test_openai_compatible_embedding_client_and_short_threshold():
    from app.services.embeddings.openai_compatible import OpenAICompatibleEmbeddings

    requests = []

    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(
            200,
            json={
                "data": [
                    {"index": 1, "embedding": [0.0, 1.0]},
                    {"index": 0, "embedding": [1.0, 0.0]},
                ]
            },
        )

    service = OpenAICompatibleEmbeddings(
        base_url="https://example.invalid/v1",
        api_key="test",
        model="embedding-test",
    )
    service.client.close()
    service.client = httpx.Client(transport=httpx.MockTransport(respond))
    try:
        vectors = service.embed_documents(["一", "二"])
    finally:
        service.close()
    assert vectors == [[1.0, 0.0], [0.0, 1.0]]
    assert requests[0]["input"] == ["一", "二"]

    class CountingEmbeddings(Embeddings):
        calls = 0

        def embed_documents(self, texts):
            self.calls += 1
            return [[1.0, 0.0] for _ in texts]

        def embed_query(self, text):
            return [1.0, 0.0]

    from app.services.documents.semantic import semantic_chunks
    from app.tasks.pipeline import make_page, text_document

    counter = CountingEmbeddings()
    document = text_document(
        {"filename": "short.md", "sha256": "a" * 64},
        [make_page(1, "短章节正文。")],
    )
    semantic_chunks(
        document,
        uuid4(),
        uuid4(),
        TinyEmbedder(),
        semantic_embedder=counter,
        openai_threshold=5000,
    )
    assert counter.calls == 0

    long_document = text_document(
        {"filename": "long.md", "sha256": "b" * 64},
        [make_page(1, "这是一个需要语义分析的完整句子。" * 400)],
    )
    semantic_chunks(
        long_document,
        uuid4(),
        uuid4(),
        TinyEmbedder(),
        semantic_embedder=counter,
        openai_threshold=5000,
    )
    assert counter.calls > 0
