import argparse
import hashlib
import json
import re
import unicodedata
from datetime import UTC, datetime
from pathlib import Path

from app.services.documents.visual_structure import (
    extract_chart,
    extract_connector_relations,
    extract_table,
)
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        while chunk := file.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_for_match(text: str) -> str:
    # 只用于比较重复内容，不覆盖原始文字。
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", "", text).casefold()


def clean_text(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\v", "\n").strip()


def extract_shapes(shapes, parent_path: str = "") -> list[dict]:
    blocks = []

    for index, shape in enumerate(shapes, start=1):
        shape_path = f"{parent_path}/{index}" if parent_path else str(index)

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            blocks.extend(extract_shapes(shape.shapes, shape_path))
            continue

        base = {
            "source": "native",
            "shape_path": shape_path,
            "shape_name": shape.name,
            "include_in_retrieval": True,
        }

        if shape.has_table:
            table_block = extract_table(shape, base)
            if table_block["text"].strip():
                blocks.append(table_block)

        elif shape.has_chart:
            blocks.append(extract_chart(shape, base))

        elif shape.has_text_frame:
            text = clean_text(shape.text_frame.text)

            if text:
                blocks.append(
                    {
                        **base,
                        "kind": "text",
                        "text": text,
                    }
                )

        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blocks.append(
                {
                    **base,
                    "kind": "image",
                    "text": "[图片]",
                    "include_in_retrieval": False,
                    "requires_review": True,
                }
            )

    return blocks


def native_signatures(blocks: list[dict]) -> set[str]:
    signatures = set()

    for block in blocks:
        text = block["text"]
        candidates = [text, *text.splitlines()]

        if block["kind"] == "table":
            for row in block["rows"]:
                candidates.extend(row)
                candidates.append(" ".join(row))

        for candidate in candidates:
            value = normalize_for_match(candidate)
            if value:
                signatures.add(value)

    return signatures


def resolve_child(directory: Path, relative_path: str) -> Path:
    path = (directory / relative_path).resolve()

    if not path.is_relative_to(directory):
        raise ValueError("清单中的结果路径超出了 OCR 目录。")

    return path


def merge(source: Path, ocr_manifest_path: Path) -> Path:
    manifest = json.loads(ocr_manifest_path.read_text(encoding="utf-8"))

    if manifest.get("stage") != "ocr_completed":
        raise ValueError("OCR 尚未全部成功，不能生成正式合并结果。")

    if manifest.get("failed_pages") != 0:
        raise ValueError("OCR 清单包含失败页，请先解决失败。")

    source_hash = sha256_file(source)

    if manifest.get("source_sha256") != source_hash:
        raise ValueError(
            "PPTX 与 OCR 来源不一致。请使用渲染时的同一份 PPTX，不能使用修改后的版本。"
        )

    presentation = Presentation(str(source))
    page_records = manifest.get("pages", [])

    if len(presentation.slides) != len(page_records) or manifest.get("page_count") != len(
        page_records
    ):
        raise ValueError("PPTX 与 OCR 清单的页数不一致。")

    if [p["page_number"] for p in page_records] != list(range(1, len(page_records) + 1)):
        raise ValueError("OCR 页码不连续或存在重复。")

    ocr_directory = ocr_manifest_path.parent
    merged_pages = []
    native_count = 0
    supplement_count = 0
    duplicate_count = 0
    review_count = 0

    for page_number, slide in enumerate(presentation.slides, start=1):
        record = page_records[page_number - 1]

        if record.get("status") != "succeeded":
            raise ValueError(f"第 {page_number} 页 OCR 未成功。")

        page_file = resolve_child(ocr_directory, record["result_file"])
        ocr = json.loads(page_file.read_text(encoding="utf-8"))

        if ocr.get("page_number") != page_number:
            raise ValueError(f"第 {page_number} 页结果文件页码不符。")

        blocks = extract_shapes(slide.shapes)
        blocks.extend(extract_connector_relations(slide.shapes))
        native_count += len(blocks)
        signatures = native_signatures(blocks)

        for line in ocr["lines"]:
            indices = line["word_indices"]

            if any(index < 0 or index >= len(ocr["words"]) for index in indices):
                raise ValueError(f"第 {page_number} 页 OCR 文字索引越界。")

            words = [ocr["words"][index] for index in indices]
            text = line["text_preview"].strip()

            if not text:
                continue

            # 采用保守的完全匹配，避免模糊匹配误删截图内容。
            duplicate = normalize_for_match(text) in signatures
            needs_review = any(word["needs_review"] for word in words)

            blocks.append(
                {
                    "source": "ocr",
                    "kind": "text",
                    "text": text,
                    "bbox_px": line["bbox_px"],
                    "coordinate_system": "image_pixels_top_left",
                    "words": words,
                    "duplicate_of_native": duplicate,
                    "include_in_retrieval": not duplicate,
                    "needs_review": needs_review,
                }
            )

            duplicate_count += int(duplicate)
            supplement_count += int(not duplicate)
            review_count += int(needs_review and not duplicate)

        # 备注不参与屏幕文字去重，避免隐藏备注抑制页面文字。
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = clean_text(frame.text) if frame else ""

            if notes:
                blocks.append(
                    {
                        "source": "native",
                        "kind": "speaker_notes",
                        "text": notes,
                        "include_in_retrieval": True,
                    }
                )
                native_count += 1

        for block_index, block in enumerate(blocks, start=1):
            block["block_id"] = f"p{page_number:04d}-b{block_index:04d}"

        merged_pages.append(
            {
                "page_number": page_number,
                "image_sha256": ocr["image_sha256"],
                "image_width_px": ocr["width_px"],
                "image_height_px": ocr["height_px"],
                "ocr_no_text_detected": ocr["no_text_detected"],
                "ocr_result_sha256": sha256_file(page_file),
                "blocks": blocks,
            }
        )

    if sha256_file(source) != source_hash:
        raise ValueError("处理期间 PPTX 发生变化，请重新执行。")

    result = {
        "schema_version": 1,
        "stage": "merged",
        "source_filename": source.name,
        "source_sha256": source_hash,
        "ocr_manifest_sha256": sha256_file(ocr_manifest_path),
        "created_at": datetime.now(UTC).isoformat(),
        "page_count": len(merged_pages),
        "deduplication": "normalized_exact_match_per_page",
        "limitations": [
            "OCR 文字中的公式不保证准确",
            "OCR 行序不保证是复杂版面的阅读顺序",
            "没有连接端点信息的箭头或纯位图图表仍不能可靠重建",
            "完全匹配去重可能保留存在识别误差的重复行",
        ],
        "statistics": {
            "native_blocks": native_count,
            "ocr_supplement_lines": supplement_count,
            "ocr_duplicate_lines": duplicate_count,
            "ocr_supplement_lines_needing_review": review_count,
        },
        "pages": merged_pages,
    }

    output = ocr_directory / "merged-document.json"
    temporary = output.with_suffix(".json.tmp")
    temporary.write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temporary.replace(output)

    print(json.dumps(result["statistics"], ensure_ascii=False, indent=2))
    print(f"合并结果：{output}")
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="合并 PPTX 原生内容和逐页 OCR 结果。")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--ocr-manifest", type=Path, required=True)
    args = parser.parse_args()

    source = args.source.expanduser().resolve()
    manifest = args.ocr_manifest.expanduser().resolve()

    if not source.is_file() or source.suffix.lower() != ".pptx":
        parser.error("--source 必须是存在的 PPTX 文件。")

    if not manifest.is_file():
        parser.error("找不到 OCR 清单。")

    merge(source, manifest)


if __name__ == "__main__":
    main()
